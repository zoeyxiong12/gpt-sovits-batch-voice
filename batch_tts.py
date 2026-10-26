# -*- coding: utf-8 -*-
"""GPT-SoVITS batch generation script: reads the xlsx directly and generates
one voice file per row from the "Japanese line" column.
Rows with an empty line are skipped; files that already exist are skipped,
so the script can be stopped and re-run safely."""
import os, sys, requests

# ============ Config ============
API_BASE   = "http://127.0.0.1:9880"
SOVITS_PTH = r"D:\GPT-SoVITS\SoVITS_weights\TODO_character.pth"
GPT_CKPT   = r"D:\GPT-SoVITS\GPT_weights\TODO_character.ckpt"
REF_WAV    = r"D:\GPT-SoVITS\DATA\character_A\slicer_opt\TODO_reference.wav"
REF_TEXT   = "TODO: transcript of the reference clip"
XLSX_PATH  = os.path.join(os.path.dirname(os.path.abspath(__file__)), "voice_list.xlsx")
OUT_DIR    = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")   # audio is saved here, in subfolders named by the "Output folder" column
CANDIDATES = 1     # versions to generate per line; 1 = one version only
# ================================

def read_rows():
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl is missing. Install it first with:\n"
                 r"D:\GPT-SoVITS\runtime\python.exe -m pip install openpyxl")
    wb = load_workbook(XLSX_PATH, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(it)]
    idx = {h: i for i, h in enumerate(header)}
    for need in ("Voice name", "Japanese line"):
        if need not in idx:
            sys.exit(f'Column "{need}" not found in the header. Current header: {header}')
    for r in it:
        def cell(name):
            i = idx.get(name)
            v = r[i] if i is not None and i < len(r) else None
            return "" if v is None else str(v).strip()
        yield cell("Voice name"), cell("Japanese line"), cell("Output folder")

def set_weights():
    if GPT_CKPT:
        r = requests.get(f"{API_BASE}/set_gpt_weights", params={"weights_path": GPT_CKPT})
        print("Switch GPT weights:", r.status_code, r.text[:100])
    r = requests.get(f"{API_BASE}/set_sovits_weights", params={"weights_path": SOVITS_PTH})
    print("Switch SoVITS weights:", r.status_code, r.text[:100])
    if r.status_code != 200:
        sys.exit("Failed to load weights. Check the paths and make sure the API is running.")

def main():
    for k, v in (('SOVITS_PTH', SOVITS_PTH), ('REF_WAV', REF_WAV), ('REF_TEXT', REF_TEXT)):
        if 'TODO' in v:
            sys.exit(f'Fill in {k} at the top of the script first.')
    set_weights()
    done = skipped = failed = 0
    for name, text, folder in read_rows():
        if not name or not text:
            skipped += 1
            continue
        out_dir = os.path.join(OUT_DIR, folder) if folder else OUT_DIR
        os.makedirs(out_dir, exist_ok=True)
        for k in range(1, CANDIDATES + 1):
            path = os.path.join(out_dir, f"{name}.wav" if CANDIDATES == 1 else f"{name}_c{k}.wav")
            if os.path.exists(path):
                continue
            r = requests.post(f"{API_BASE}/tts", json={
                "text": text, "text_lang": "ja",
                "ref_audio_path": REF_WAV,
                "prompt_text": REF_TEXT, "prompt_lang": "ja",
                "top_k": 15, "top_p": 1, "temperature": 1,
                "text_split_method": "cut0",
                "speed_factor": 1.0, "seed": -1,
                "media_type": "wav", "streaming_mode": False,
            })
            if r.status_code == 200:
                with open(path, "wb") as g:
                    g.write(r.content)
                done += 1
                print("OK  ", path)
            else:
                failed += 1
                print("FAIL", name, r.text[:200])
    print(f"\nDone {done}, skipped (empty line) {skipped}, failed {failed}")

if __name__ == "__main__":
    main()
